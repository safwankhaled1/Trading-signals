from __future__ import annotations

import csv
from datetime import datetime, timedelta, timezone
from pathlib import Path

from .store import Store


def build_report(path, start, end, channel=None, scope="all", utc_offset=0, local=False, account=None):
    store = Store(path)
    try:
        if account:
            store.account = account
        zone = datetime.now().astimezone().tzinfo if local else timezone(timedelta(hours=float(utc_offset)))
        def day_boundary(text):
            dt = datetime.fromisoformat(text)
            if local:
                return dt.astimezone().timestamp()
            return dt.replace(tzinfo=zone).timestamp()
        left, right = day_boundary(start), day_boundary(end) + 86400
        if right <= left:
            raise ValueError("تاريخ نهاية التقرير يسبق البداية")
        signals = {s["id"]: s for s in store.signals()}
        deals = [d for d in store.deals() if left <= d["time"] < right
                 and (channel is None or str(d.get("channel")) == str(channel))
                 and (scope == "all" or bool(d.get("manual")) == (scope == "manual"))]
        grouped = {}
        daily = {}
        rows = []
        for d in sorted(deals, key=lambda d: d["time"]):
            signal = signals.get(d["signal"], {})
            grouped[d["signal"]] = grouped.get(d["signal"], 0) + d["net"]
            stamp = datetime.fromtimestamp(d["time"]).astimezone() if local else datetime.fromtimestamp(d["time"], zone)
            day = stamp.strftime("%Y-%m-%d")
            daily[day] = daily.get(day, 0) + d["net"]
            rows.append({"time": stamp.strftime("%Y-%m-%d %H:%M:%S"), "channel": d.get("channel_name", ""), "signal": d["signal"],
                         "position": d["position_id"], "type": "دخول" if d["entry_type"] == 0 else "إغلاق",
                         "volume": d["volume"], "price": d["price"], "profit": round(d["profit"], 2),
                         "commission": d["commission"], "swap": d["swap"], "fee": d["fee"], "net": round(d["net"], 2),
                         "state": signal.get("state", "")})
        # Win rate uses fully closed signals, with all their lifecycle fees/PnL, not partial closes.
        closed = []
        all_deals = store.deals()
        for identifier in grouped:
            s = signals.get(identifier, {})
            if s.get("state") == "closed" and left <= s.get("closed", 0) < right:
                closed.append(sum(d["net"] for d in all_deals if d["signal"] == identifier))
        return {"account": account, "start": start, "end": end, "net": round(sum(d["net"] for d in deals), 2), "count": len(grouped),
                "closed_count": len(closed), "win_rate": round(100 * sum(v > 0 for v in closed) / len(closed), 1) if closed else 0,
                "best": round(max(closed, default=0), 2), "worst": round(min(closed, default=0), 2),
                "daily": [{"day": day, "net": round(value, 2)} for day, value in sorted(daily.items())], "rows": rows,
                "timezone": "محلي" if local else f"UTC{float(utc_offset):+g}"}
    finally:
        store.close()


HEADERS = {"time": "الوقت", "channel": "القناة", "signal": "الإشارة", "position": "رقم الصفقة", "type": "النوع",
           "volume": "اللوت", "price": "السعر", "profit": "الربح", "commission": "العمولة", "swap": "السواب", "fee": "الرسوم", "net": "الصافي", "state": "الحالة"}


def export_report(report, path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "تفاصيل التداول"
        ws.sheet_view.rightToLeft = True
        ws.append(list(HEADERS.values()))
        for row in report["rows"]:
            ws.append([row.get(k, "") for k in HEADERS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173D36")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(32, max(14, max(len(str(c.value or "")) for c in col) + 2))
            for cell in col:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
        summary = wb.create_sheet("الملخص")
        summary.sheet_view.rightToLeft = True
        summary.append(["الحساب", report.get("account") or "محاكاة"])
        for k, title in [("start", "من"), ("end", "إلى"), ("net", "صافي الربح"), ("count", "إشارات لها عمليات"), ("closed_count", "صفقات مغلقة"), ("win_rate", "نسبة النجاح %"), ("timezone", "توقيت التقرير")]:
            summary.append([title, report[k]])
        summary.column_dimensions["A"].width = 26
        summary.column_dimensions["B"].width = 24
        wb.save(path)
    else:
        with path.open("w", newline="", encoding="utf-8-sig") as output:
            writer = csv.writer(output)
            writer.writerow(HEADERS.values())
            for row in report["rows"]:
                values = [row.get(k, "") for k in HEADERS]
                writer.writerow(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v for v in values])
    return str(path)
